"""Un administrateur simple ne touche pas au super-administrateur : ni mot de passe, ni blocage,
ni fiche, ni actions tracées dans le journal d'audit (liste et export). Tests sur les vraies routes."""
from io import BytesIO

from django.contrib.auth import get_user_model
from django.urls import reverse
from openpyxl import load_workbook
from rest_framework.test import APITestCase

from jobs.models import AuditLog

User = get_user_model()


class ProtectionSuperAdminTest(APITestCase):
    def setUp(self):
        self.boss = User.objects.create_superuser(
            username='boss', email='boss@test.com', password='MotDePasse-Boss-1', role='ADMIN')
        self.admin = User.objects.create_user(
            username='adm', email='adm@test.com', password='x-Pass-1234', role='ADMIN', is_staff=True)
        AuditLog.objects.create(admin=self.boss, action='AUTRE', detail='action-du-boss')
        AuditLog.objects.create(admin=self.admin, action='AUTRE', detail='action-de-ladmin')

    def _details(self, response):
        return [ligne['detail'] for ligne in response.data['results']]

    def test_admin_simple_ne_change_pas_le_mot_de_passe_du_superadmin(self):
        self.client.force_authenticate(user=self.admin)
        r = self.client.patch(
            reverse('admin-compte-admin-detail', kwargs={'admin_id': self.boss.id}),
            {'password': 'Nouveau-Pass-1234'}, format='json')
        self.assertEqual(r.status_code, 403)
        self.boss.refresh_from_db()
        self.assertTrue(self.boss.check_password('MotDePasse-Boss-1'))

    def test_admin_simple_ne_bloque_pas_le_superadmin(self):
        self.client.force_authenticate(user=self.admin)
        r = self.client.patch(reverse('admin-user-moderer', kwargs={'user_id': self.boss.id}))
        self.assertEqual(r.status_code, 403)
        self.boss.refresh_from_db()
        self.assertTrue(self.boss.is_active)

    def test_admin_simple_ne_voit_pas_la_fiche_du_superadmin(self):
        self.client.force_authenticate(user=self.admin)
        r = self.client.get(reverse('admin-dossier-admin', kwargs={'user_id': self.boss.id}))
        self.assertEqual(r.status_code, 404)

    def test_journal_cache_les_actions_du_superadmin_a_ladmin_simple(self):
        self.client.force_authenticate(user=self.admin)
        details = self._details(self.client.get(reverse('admin-audit-logs')))
        self.assertIn('action-de-ladmin', details)
        self.assertNotIn('action-du-boss', details)

    def test_recherche_du_journal_ne_revele_pas_le_superadmin(self):
        self.client.force_authenticate(user=self.admin)
        r = self.client.get(reverse('admin-audit-logs'), {'search': 'boss@test.com'})
        self.assertEqual(self._details(r), [])

    def test_export_du_journal_exclut_le_superadmin_pour_ladmin_simple(self):
        self.client.force_authenticate(user=self.admin)
        r = self.client.get(reverse('export-audit-logs'))
        self.assertEqual(r.status_code, 200)
        cellules = [str(c.value) for ligne in load_workbook(BytesIO(r.content)).active.iter_rows() for c in ligne]
        self.assertNotIn('action-du-boss', cellules)
        self.assertIn('action-de-ladmin', cellules)

    def test_liste_des_comptes_admins_marque_le_superadmin_non_modifiable(self):
        url = reverse('admin-comptes-admins')
        self.client.force_authenticate(user=self.admin)
        par_email = {ligne['email']: ligne['modifiable'] for ligne in self.client.get(url).data}
        self.assertFalse(par_email['boss@test.com'])
        self.assertTrue(par_email['adm@test.com'])
        self.client.force_authenticate(user=self.boss)
        par_email = {ligne['email']: ligne['modifiable'] for ligne in self.client.get(url).data}
        self.assertTrue(par_email['boss@test.com'])

    def test_superadmin_garde_la_vue_complete(self):
        self.client.force_authenticate(user=self.boss)
        details = self._details(self.client.get(reverse('admin-audit-logs')))
        self.assertIn('action-du-boss', details)
        self.assertEqual(
            self.client.get(reverse('admin-dossier-admin', kwargs={'user_id': self.boss.id})).status_code, 200)
