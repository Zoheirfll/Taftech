# test_api_admin_kpi.py — page admin « KPI » (route admin/kpi/, calcul dans stats_admin/kpi.py)
import datetime
import io

from django.contrib.auth import get_user_model
from django.test import TestCase
from django.utils import timezone
from openpyxl import load_workbook
from rest_framework import status
from rest_framework.test import APIClient

from jobs.models import (
    AuditLog, Candidature, DemandeAjout, OffreEmploi, PaiementAbonnement, PaiementCreditPack,
    ParsingCV, ProfilCandidat, ProfilEntreprise, UtilisationIA,
)

User = get_user_model()
URL = '/api/jobs/admin/kpi/'


class AdminKpiTests(TestCase):
    def setUp(self):
        self.client = APIClient()
        self.admin = User.objects.create_user(
            username='adm', email='adm@test.dz', password='Pass1234!', role='ADMIN', is_staff=True)
        self.candidat = User.objects.create_user(
            username='cand', email='cand@test.dz', password='Pass1234!', role='CANDIDAT')
        self.recruteur = User.objects.create_user(
            username='rec', email='rec@test.dz', password='Pass1234!', role='RECRUTEUR')

    def _get(self, **params):
        self.client.force_authenticate(user=self.admin)
        return self.client.get(URL, params)

    def test_reserve_aux_admins(self):
        self.client.force_authenticate(user=None)
        self.assertEqual(self.client.get(URL).status_code, status.HTTP_401_UNAUTHORIZED)
        for user in (self.candidat, self.recruteur):
            with self.subTest(role=user.role):
                self.client.force_authenticate(user=user)
                self.assertEqual(self.client.get(URL).status_code, status.HTTP_403_FORBIDDEN)

    def test_structure_et_aucune_donnee_personnelle(self):
        reponse = self._get()
        self.assertEqual(reponse.status_code, status.HTTP_200_OK)
        data = reponse.json()
        self.assertEqual(set(data), {'periode', 'croissance', 'activite', 'revenus', 'ia', 'a_traiter'})
        self.assertNotIn('cand@test.dz', reponse.content.decode())
        self.assertNotIn('rec@test.dz', reponse.content.decode())

    def test_inscriptions_et_variation(self):
        # Un candidat inscrit pendant la période (setUp) ; rien sur la période précédente -> pas de variation.
        data = self._get().json()
        inscrits = data['croissance']['inscriptions_candidats']
        self.assertEqual(inscrits['valeur'], 1)
        self.assertEqual(inscrits['precedente'], 0)
        self.assertIsNone(inscrits['variation_pct'])

    def test_ia_compte_les_appels_acceptes_et_pas_les_rembourses(self):
        UtilisationIA.objects.create(user=self.candidat, fonction='PARSER_CV', tokens_entree=100, tokens_sortie=50)
        UtilisationIA.objects.create(user=self.candidat, fonction='PARSER_CV', tokens_entree=999, statut='REMBOURSE')
        ia = self._get().json()['ia']
        self.assertEqual(ia['appels']['valeur'], 1)
        self.assertEqual(ia['tokens']['valeur'], 150)

    def test_taux_echec_cv_none_sans_parsing_puis_calcule(self):
        self.assertIsNone(self._get().json()['ia']['taux_echec_cv_pct']['valeur'])
        ParsingCV.objects.create(user=self.candidat, succes=True)
        ParsingCV.objects.create(user=self.candidat, succes=False, code_erreur='X')
        ia = self._get().json()['ia']
        self.assertEqual(ia['cv_analyses']['valeur'], 2)
        self.assertEqual(ia['taux_echec_cv_pct']['valeur'], 50.0)

    def test_a_traiter_compte_les_demandes_ouvertes(self):
        DemandeAjout.objects.create(
            type='METIER', libelle='Cariste', cle_recherche='cariste', origine='CANDIDAT')
        self.assertEqual(self._get().json()['a_traiter']['demandes_ajout'], 1)

    def test_periode_invalide_refusee(self):
        for params in ({'date_debut': 'hier', 'date_fin': '2026-01-01'}, {'date_debut': '2026-01-01'}):
            with self.subTest(params=params):
                self.assertEqual(self._get(**params).status_code, status.HTTP_400_BAD_REQUEST)

    # --- Détail par thème (admin/kpi/<theme>/) et export ---
    def _detail(self, theme, **params):
        self.client.force_authenticate(user=self.admin)
        return self.client.get(f'{URL}{theme}/', params)

    def _entreprise_avec_offre(self):
        entreprise = ProfilEntreprise.objects.create(
            user=self.recruteur, nom_entreprise='TechCorp', secteur_activite='IT',
            registre_commerce='RC1', wilaya_siege='31 - Oran', est_approuvee=True)
        offre = OffreEmploi.objects.create(
            entreprise=entreprise, titre='Dev', wilaya='31 - Oran', type_contrat='CDI',
            experience_requise='DEBUTANT', statut_moderation='APPROUVEE')
        return entreprise, offre

    def test_detail_reserve_aux_admins_et_theme_inconnu(self):
        self.client.force_authenticate(user=None)
        self.assertEqual(self.client.get(f'{URL}ia/').status_code, status.HTTP_401_UNAUTHORIZED)
        for user in (self.candidat, self.recruteur):
            self.client.force_authenticate(user=user)
            for url in (f'{URL}ia/', f'{URL}export/'):
                with self.subTest(role=user.role, url=url):
                    self.assertEqual(self.client.get(url).status_code, status.HTTP_403_FORBIDDEN)
        self.assertEqual(self._detail('inconnu').status_code, status.HTTP_404_NOT_FOUND)

    def test_detail_structure_commune_aux_quatre_themes(self):
        for theme in ('croissance', 'activite', 'revenus', 'ia'):
            with self.subTest(theme=theme):
                reponse = self._detail(theme)
                self.assertEqual(reponse.status_code, status.HTTP_200_OK)
                data = reponse.json()
                self.assertEqual(data['theme'], theme)
                self.assertTrue(data['evolution']['series'])
                self.assertTrue(data['evolution']['lignes'])
                for cle in ('repartitions', 'classements', 'entonnoirs'):
                    self.assertIn(cle, data)
                for serie in data['evolution']['series']:
                    self.assertIn(serie['cle'], data['evolution']['lignes'][0])

    def test_detail_periode_invalide_refusee(self):
        self.assertEqual(self._detail('ia', date_debut='hier', date_fin='2026-01-01').status_code,
                         status.HTTP_400_BAD_REQUEST)

    def test_detail_croissance_repartition_par_wilaya_et_entonnoir(self):
        ProfilCandidat.objects.create(user=self.candidat, wilaya='31 - Oran')
        data = self._detail('croissance').json()
        wilayas = next(r for r in data['repartitions'] if r['cle'] == 'candidats_wilaya')
        self.assertEqual(wilayas['lignes'], [{'libelle': '31 - Oran', 'nb': 1}])
        parcours = next(e for e in data['entonnoirs'] if e['cle'] == 'candidats')
        self.assertEqual([e['nb'] for e in parcours['etapes']], [1, 1, 0, 0])

    def test_detail_activite_classements_et_repartitions(self):
        entreprise, offre = self._entreprise_avec_offre()
        Candidature.objects.create(offre=offre, candidat=self.candidat, statut='RETENU', score_matching=70)
        data = self._detail('activite').json()
        top_offres = next(c for c in data['classements'] if c['cle'] == 'top_offres')
        self.assertEqual(top_offres['lignes'], [['Dev', 'TechCorp', 1]])
        contrats = next(r for r in data['repartitions'] if r['cle'] == 'offres_contrat')
        self.assertEqual(contrats['lignes'][0]['nb'], 1)
        entonnoir = next(e for e in data['entonnoirs'] if e['cle'] == 'recrutement')
        self.assertEqual(entonnoir['etapes'][-1]['nb'], 1)
        # Aucune donnée personnelle de candidat dans la réponse.
        self.assertNotIn('cand@test.dz', str(data))

    def test_detail_revenus_classe_les_entreprises(self):
        entreprise, _ = self._entreprise_avec_offre()
        PaiementAbonnement.objects.create(entreprise=entreprise, palier_nom='PRO', montant_da=15000)
        PaiementCreditPack.objects.create(entreprise=entreprise, pack_nom='Pack 10', credits=10, montant_da=5000)
        data = self._detail('revenus').json()
        top = next(c for c in data['classements'] if c['cle'] == 'top_revenu')
        self.assertEqual(top['lignes'], [['TechCorp', 15000, 5000, 20000]])
        palier = next(r for r in data['repartitions'] if r['cle'] == 'revenu_palier')
        self.assertEqual(palier['lignes'], [{'libelle': 'Pro', 'nb': 15000}])

    def test_detail_ia_repartit_par_fonction_sans_les_rembourses(self):
        UtilisationIA.objects.create(user=self.candidat, fonction='PARSER_CV', tokens_entree=100, tokens_sortie=50)
        UtilisationIA.objects.create(user=self.candidat, fonction='GENERATION_OFFRE', statut='REMBOURSE')
        ParsingCV.objects.create(user=self.candidat, succes=False, code_erreur='PDF_ILLISIBLE')
        data = self._detail('ia').json()
        fonctions = next(r for r in data['repartitions'] if r['cle'] == 'appels_fonction')
        self.assertEqual(len(fonctions['lignes']), 1)
        tokens = next(r for r in data['repartitions'] if r['cle'] == 'tokens_fonction')
        self.assertEqual(tokens['lignes'][0]['nb'], 150)
        erreurs = next(r for r in data['repartitions'] if r['cle'] == 'cv_erreurs')
        self.assertEqual(erreurs['lignes'], [{'libelle': 'PDF_ILLISIBLE', 'nb': 1}])

    def test_export_un_onglet_par_theme_et_trace_audit(self):
        self.client.force_authenticate(user=self.admin)
        reponse = self.client.get(f'{URL}export/')
        self.assertEqual(reponse.status_code, status.HTTP_200_OK)
        classeur = load_workbook(io.BytesIO(reponse.content))
        self.assertEqual(classeur.sheetnames, ['Croissance', 'Offres et candidatures', 'Revenus et abonnements', 'Qualité et IA'])
        self.assertEqual([c.value for c in classeur['Croissance'][1]], ['Rubrique', 'Élément', 'Valeur', 'Détail'])
        self.assertTrue(AuditLog.objects.filter(admin=self.admin, action='AUTRE', detail__startswith='Export KPI').exists())

    def test_periode_explicite_exclut_les_inscriptions_hors_periode(self):
        hier = timezone.localdate() - datetime.timedelta(days=400)
        data = self._get(date_debut=hier.isoformat(), date_fin=(hier + datetime.timedelta(days=5)).isoformat()).json()
        self.assertEqual(data['croissance']['inscriptions_candidats']['valeur'], 0)
