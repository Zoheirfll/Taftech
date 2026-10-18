# test_api_admin_statistiques.py — page admin « Statistiques » (période, variations, entonnoir, export)
import datetime
import io

from django.contrib.auth import get_user_model
from django.test import SimpleTestCase, TestCase
from django.utils import timezone
from openpyxl import load_workbook
from rest_framework import status
from rest_framework.test import APIClient

from jobs.models import (
    AbonnementEntreprise, AuditLog, Candidature, OffreEmploi, PaiementAbonnement, Palier,
    ProfilEntreprise,
)
from jobs.stats_admin.periode import (
    DUREE_MAX_JOURS, PeriodeInvalide, lire_periode, mois_de_la_periode,
    periode_par_defaut, periode_precedente, variation_pct,
)

User = get_user_model()
URL = '/api/jobs/admin/stats-periode/'
URL_EXPORT = '/api/jobs/admin/stats-periode/export/'


class PeriodeTests(SimpleTestCase):
    def test_defaut_365_derniers_jours(self):
        debut, fin = lire_periode({})
        self.assertEqual((fin - debut).days, 365)
        self.assertEqual((debut, fin), periode_par_defaut())

    def test_periode_explicite(self):
        debut, fin = lire_periode({'date_debut': '2026-01-01', 'date_fin': '2026-03-31'})
        self.assertEqual((debut, fin), (datetime.date(2026, 1, 1), datetime.date(2026, 3, 31)))

    def test_refuse_format_ordre_duree_et_parametre_seul(self):
        cas = [
            {'date_debut': 'hier', 'date_fin': '2026-03-31'},
            {'date_debut': '2026-04-01', 'date_fin': '2026-03-31'},
            {'date_debut': '2026-01-01'},
            {'date_debut': '2000-01-01', 'date_fin': '2026-03-31'},
        ]
        for params in cas:
            with self.subTest(params=params), self.assertRaises(PeriodeInvalide):
                lire_periode(params)
        self.assertGreater(DUREE_MAX_JOURS, 365 * 5)

    def test_periode_precedente_meme_duree_juste_avant(self):
        debut, fin = datetime.date(2026, 3, 1), datetime.date(2026, 3, 31)
        prec_debut, prec_fin = periode_precedente(debut, fin)
        self.assertEqual(prec_fin, datetime.date(2026, 2, 28))
        self.assertEqual((prec_fin - prec_debut).days, (fin - debut).days)

    def test_variation_none_si_precedente_nulle(self):
        self.assertIsNone(variation_pct(5, 0))
        self.assertEqual(variation_pct(15, 10), 50.0)
        self.assertEqual(variation_pct(5, 10), -50.0)

    def test_mois_vides_inclus(self):
        mois = mois_de_la_periode(datetime.date(2026, 1, 15), datetime.date(2026, 4, 2))
        self.assertEqual([m.month for m in mois], [1, 2, 3, 4])


class AdminStatistiquesAPITests(TestCase):
    def setUp(self):
        self.client = APIClient()
        self.admin = User.objects.create_user(
            username='adm', email='adm@test.dz', password='Pass1234!', role='ADMIN', is_staff=True)
        self.recruteur = User.objects.create_user(
            username='rec', email='rec@test.dz', password='Pass1234!', role='RECRUTEUR')
        self.candidat = User.objects.create_user(
            username='cand', email='cand@test.dz', password='Pass1234!', role='CANDIDAT')
        self.entreprise = ProfilEntreprise.objects.create(
            user=self.recruteur, nom_entreprise='TechCorp', secteur_activite='IT',
            registre_commerce='RC1', wilaya_siege='31 - Oran', est_approuvee=True)
        self.offre = self._offre('Dev')

    def _offre(self, titre):
        return OffreEmploi.objects.create(
            entreprise=self.entreprise, titre=titre, wilaya='31 - Oran',
            type_contrat='CDI', experience_requise='DEBUTANT', statut_moderation='APPROUVEE')

    def _candidature(self, statut='RECUE', score=None, offre=None):
        return Candidature.objects.create(
            offre=offre or self.offre, candidat=self.candidat, statut=statut, score_matching=score)

    def _get(self, params=None, url=URL):
        self.client.force_authenticate(user=self.admin)
        return self.client.get(url, params or {})

    # --- Permissions (vraie route) ---
    def test_refuse_anonyme_candidat_et_recruteur(self):
        self.assertIn(self.client.get(URL).status_code, (401, 403))
        for user in (self.candidat, self.recruteur):
            for url in (URL, URL_EXPORT):
                self.client.force_authenticate(user=user)
                self.assertEqual(self.client.get(url).status_code, status.HTTP_403_FORBIDDEN)

    def test_periode_invalide_renvoie_400_avec_message(self):
        rep = self._get({'date_debut': '2026-05-01', 'date_fin': '2026-01-01'})
        self.assertEqual(rep.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn('error', rep.data)

    # --- Contenu ---
    def test_structure_et_defaut(self):
        data = self._get().data
        self.assertEqual(
            set(data),
            {'periode', 'entreprise', 'entreprises', 'indicateurs', 'evolution', 'entonnoir', 'offres',
             'matching', 'abonnements'})
        self.assertIsNone(data['entreprise'])
        self.assertEqual(data['entreprises'], [{'id': self.entreprise.id, 'nom': 'TechCorp'}])
        self.assertEqual(data['indicateurs']['offres']['valeur'], 1)

    def test_indicateurs_et_variation_vs_periode_precedente(self):
        self._candidature()
        ancienne = self._candidature()
        il_y_a_40_jours = timezone.now() - datetime.timedelta(days=40)
        Candidature.objects.filter(pk=ancienne.pk).update(date_postulation=il_y_a_40_jours)
        debut = (timezone.localdate() - datetime.timedelta(days=29)).isoformat()
        fin = timezone.localdate().isoformat()
        ind = self._get({'date_debut': debut, 'date_fin': fin}).data['indicateurs']['candidatures']
        self.assertEqual(ind['valeur'], 1)
        self.assertEqual(ind['precedente'], 1)
        self.assertEqual(ind['variation_pct'], 0.0)

    def test_variation_none_quand_periode_precedente_vide(self):
        self._candidature()
        ind = self._get().data['indicateurs']['candidatures']
        self.assertEqual(ind['valeur'], 1)
        self.assertIsNone(ind['variation_pct'])

    def test_evolution_inclut_les_mois_vides(self):
        debut = (timezone.localdate() - datetime.timedelta(days=100)).isoformat()
        evo = self._get({'date_debut': debut, 'date_fin': timezone.localdate().isoformat()}).data['evolution']
        self.assertGreaterEqual(len(evo), 4)
        self.assertEqual(sum(m['offres'] for m in evo), 1)

    def test_entonnoir_cumulatif_et_taux_de_passage(self):
        for statut in ('RECUE', 'EN_COURS', 'PRESELECTION', 'RETENU', 'REFUSE'):
            self._candidature(statut)
        ent = self._get().data['entonnoir']
        nb = {e['cle']: e['nb'] for e in ent['etapes']}
        self.assertEqual(nb, {'deposees': 5, 'examinees': 3, 'preselectionnees': 2, 'entretiens': 1, 'retenues': 1})
        self.assertEqual(ent['refusees'], 1)
        self.assertEqual(ent['etapes'][1]['taux_passage'], 60.0)

    def test_entonnoir_vide_sans_division_par_zero(self):
        ent = self._get().data['entonnoir']
        self.assertTrue(all(e['nb'] == 0 and e['taux_passage'] is None for e in ent['etapes']))

    def test_offres_sans_candidature_et_expirees(self):
        expiree = self._offre('Expirée')
        OffreEmploi.objects.filter(pk=expiree.pk).update(date_expiration=timezone.localdate() - datetime.timedelta(days=2))
        self._candidature()
        offres = self._get().data['offres']
        self.assertEqual(offres['offres_approuvees'], 2)
        self.assertEqual(offres['sans_candidature'], 1)
        self.assertEqual(offres['expirees_sans_candidature'], 1)
        self.assertEqual(offres['candidatures_par_offre'], 0.5)

    def test_distribution_matching(self):
        self._candidature(score=10)
        self._candidature(score=85)
        self._candidature(score=100)
        self._candidature(score=None)
        matching = self._get().data['matching']
        nb = {t['tranche']: t['nb'] for t in matching['distribution']}
        self.assertEqual((nb['0-20'], nb['80-100']), (1, 2))
        self.assertEqual(matching['moyenne'], 65.0)

    def test_revenu_par_palier_et_mensuel(self):
        PaiementAbonnement.objects.create(entreprise=self.entreprise, palier_nom='PRO', montant_da=5000)
        PaiementAbonnement.objects.create(entreprise=self.entreprise, palier_nom='PRO', montant_da=5000)
        data = self._get().data
        self.assertEqual(data['indicateurs']['revenu_da']['valeur'], 10000)
        self.assertEqual(data['abonnements']['revenu_par_palier'][0]['total_da'], 10000)
        self.assertEqual(sum(m['revenu_da'] for m in data['evolution']), 10000)

    def test_aucune_donnee_personnelle_exposee(self):
        self._candidature()
        data = self._get().data
        brut = self._get().content.decode()
        for secret in ('cand@test.dz', 'rec@test.dz', 'adm@test.dz', 'RC1', '31 - Oran'):
            self.assertNotIn(secret, brut)
        # Le nom d'entreprise n'est exposé que dans la liste du sélecteur (identifiant + nom).
        sans_selecteur = {cle: valeur for cle, valeur in data.items() if cle != 'entreprises'}
        self.assertNotIn('TechCorp', str(sans_selecteur))
        self.assertEqual(set(data['entreprises'][0]), {'id', 'nom'})

    # --- Export ---
    def test_export_xlsx_onglets_et_audit(self):
        self._candidature('RETENU', score=70)
        rep = self._get(url=URL_EXPORT)
        self.assertEqual(rep.status_code, 200)
        self.assertIn('spreadsheetml', rep['Content-Type'])
        classeur = load_workbook(io.BytesIO(rep.content))
        self.assertEqual(
            classeur.sheetnames,
            ['Indicateurs', 'Évolution mensuelle', 'Entonnoir', 'Offres', 'Matching', 'Abonnements'])
        self.assertTrue(AuditLog.objects.filter(admin=self.admin, action='AUTRE', detail__startswith='Export statistiques').exists())

    def test_export_refuse_periode_invalide_sans_audit(self):
        rep = self._get({'date_debut': 'x', 'date_fin': 'y'}, url=URL_EXPORT)
        self.assertEqual(rep.status_code, 400)
        self.assertFalse(AuditLog.objects.filter(detail__startswith='Export statistiques').exists())

    # --- Filtre par entreprise ---
    def _autre_entreprise(self):
        autre_user = User.objects.create_user(
            username='rec2', email='rec2@test.dz', password='Pass1234!', role='RECRUTEUR')
        autre = ProfilEntreprise.objects.create(
            user=autre_user, nom_entreprise='AutreCorp', secteur_activite='IT',
            registre_commerce='RC2', wilaya_siege='16 - Alger', est_approuvee=True)
        return autre, self._offre_de(autre, 'Compta')

    def _offre_de(self, entreprise, titre):
        return OffreEmploi.objects.create(
            entreprise=entreprise, titre=titre, wilaya='16 - Alger',
            type_contrat='CDI', experience_requise='DEBUTANT', statut_moderation='APPROUVEE')

    def test_filtre_entreprise_restreint_offres_candidatures_et_revenu(self):
        autre, offre_autre = self._autre_entreprise()
        self._candidature('RETENU', score=80)
        self._candidature('RECUE', offre=offre_autre)
        self._candidature('RECUE', offre=offre_autre)
        PaiementAbonnement.objects.create(entreprise=autre, palier_nom='PRO', montant_da=7000)
        data = self._get({'entreprise_id': self.entreprise.id}).data
        ind = data['indicateurs']
        self.assertEqual((ind['offres']['valeur'], ind['candidatures']['valeur']), (1, 1))
        self.assertEqual(ind['recrutements']['valeur'], 1)
        self.assertEqual(ind['revenu_da']['valeur'], 0)
        self.assertEqual(data['entonnoir']['etapes'][0]['nb'], 1)
        self.assertEqual(data['matching']['moyenne'], 80.0)
        self.assertEqual(data['entreprise'], {'id': self.entreprise.id, 'nom': 'TechCorp'})
        # Sans filtre : tout le monde.
        tout = self._get().data['indicateurs']
        self.assertEqual((tout['offres']['valeur'], tout['candidatures']['valeur']), (2, 3))
        self.assertEqual(tout['revenu_da']['valeur'], 7000)

    def test_filtre_entreprise_masque_les_inscriptions(self):
        data = self._get({'entreprise_id': self.entreprise.id}).data
        self.assertNotIn('inscriptions_candidats', data['indicateurs'])
        self.assertNotIn('inscriptions_recruteurs', data['indicateurs'])
        self.assertTrue(all('inscriptions_candidats' not in m for m in data['evolution']))
        self.assertIn('inscriptions_candidats', self._get().data['indicateurs'])

    def test_filtre_entreprise_invalide_ou_inconnue_renvoie_400(self):
        for valeur in ('abc', '-3', '0', '999999'):
            for url in (URL, URL_EXPORT):
                with self.subTest(valeur=valeur, url=url):
                    rep = self._get({'entreprise_id': valeur}, url=url)
                    self.assertEqual(rep.status_code, status.HTTP_400_BAD_REQUEST)
                    self.assertIn('error', rep.data)

    def test_export_avec_filtre_entreprise_indique_le_perimetre(self):
        rep = self._get({'entreprise_id': self.entreprise.id}, url=URL_EXPORT)
        self.assertEqual(rep.status_code, 200)
        self.assertIn(f'_entreprise-{self.entreprise.id}', rep['Content-Disposition'])
        feuille = load_workbook(io.BytesIO(rep.content))['Indicateurs']
        lignes = [[c.value for c in ligne] for ligne in feuille.iter_rows()]
        self.assertIn('TechCorp', [ligne[1] for ligne in lignes if ligne[0] == 'Périmètre'])
        self.assertTrue(AuditLog.objects.filter(detail__contains=f'entreprise {self.entreprise.id}').exists())


class AdminEntreprisesFiltreExpirationTests(TestCase):
    """Lien de la bannière « abonnements qui expirent » du tableau de bord."""
    URL = '/api/jobs/admin/entreprises/'

    def setUp(self):
        self.client = APIClient()
        self.admin = User.objects.create_user(
            username='adm', email='adm@test.dz', password='Pass1234!', role='ADMIN', is_staff=True)
        self.client.force_authenticate(user=self.admin)
        self.palier, _ = Palier.objects.get_or_create(nom='PRO')

    def _entreprise(self, nom, jours=None):
        user = User.objects.create_user(
            username=nom, email=f'{nom}@test.dz', password='Pass1234!', role='RECRUTEUR')
        entreprise = ProfilEntreprise.objects.create(
            user=user, nom_entreprise=nom, secteur_activite='IT',
            registre_commerce=f'RC-{nom}', wilaya_siege='31 - Oran', est_approuvee=True)
        if jours is not None:
            AbonnementEntreprise.objects.create(
                entreprise=entreprise, palier=self.palier,
                date_expiration=timezone.now() + datetime.timedelta(days=jours))
        return entreprise

    def _noms(self, params=None):
        rep = self.client.get(self.URL, params or {})
        self.assertEqual(rep.status_code, 200)
        return {e['nom_entreprise'] for e in rep.data['results']}

    def test_filtre_garde_les_abonnements_qui_expirent_sous_7_jours(self):
        self._entreprise('Bientot', jours=3)
        self._entreprise('Loin', jours=20)
        self._entreprise('Expire', jours=-2)
        self._entreprise('Gratuite')
        self.assertEqual(self._noms({'abonnement_expire_sous': 7}), {'Bientot'})
        self.assertEqual(self._noms({'abonnement_expire_sous': 30}), {'Bientot', 'Loin'})
        self.assertEqual(len(self._noms()), 4)

    def test_valeur_invalide_ignoree(self):
        self._entreprise('Loin', jours=20)
        for valeur in ('abc', '-1', '0', '9999'):
            self.assertEqual(self._noms({'abonnement_expire_sous': valeur}), {'Loin'})

    def test_expose_la_date_d_expiration(self):
        self._entreprise('Bientot', jours=3)
        self._entreprise('Gratuite')
        rep = self.client.get(self.URL)
        par_nom = {e['nom_entreprise']: e for e in rep.data['results']}
        self.assertIsNotNone(par_nom['Bientot']['abonnement_expire_at'])
        self.assertIsNone(par_nom['Gratuite']['abonnement_expire_at'])
