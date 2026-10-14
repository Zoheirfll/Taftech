"""Export .xlsx des statistiques/graphiques recruteur (correction client n°16).

L'ancien CSV fait à la main côté client (séparateur « ; » + BOM) s'ouvrait avec toutes les
données dans une seule colonne sur smartphone ; l'export est désormais un vrai classeur Excel.
"""
from io import BytesIO

from django.contrib.auth import get_user_model
from django.core.cache import cache
from django.urls import reverse
from openpyxl import load_workbook
from rest_framework import status
from rest_framework.test import APITestCase

User = get_user_model()

XLSX_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


class ExportGraphiqueExcelTests(APITestCase):
    def setUp(self):
        cache.clear()  # le throttle d'écriture s'appuie sur le cache
        self.rh = User.objects.create_user(username="rh_xl", email="rh_xl@test.com", role="RECRUTEUR")
        self.url = reverse("export-graphique-excel")
        self.payload = {
            "titre": "statistiques-recrutement",
            "colonnes": ["Offre", "Statut offre", "Candidatures", "Taux de conversion (%)"],
            "lignes": [["Data analyst", "Active", 12, 25], ["Comptable", "Clôturée", 4, 0]],
        }

    def test_anonyme_refuse(self):
        response = self.client.post(self.url, self.payload, format="json")
        self.assertIn(response.status_code, (status.HTTP_401_UNAUTHORIZED, status.HTTP_403_FORBIDDEN))

    def test_genere_un_vrai_classeur_xlsx_en_colonnes_separees(self):
        self.client.force_authenticate(user=self.rh)
        response = self.client.post(self.url, self.payload, format="json")

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response["Content-Type"], XLSX_TYPE)
        self.assertIn("statistiques-recrutement.xlsx", response["Content-Disposition"])

        ws = load_workbook(BytesIO(response.content)).active
        lignes = [[c.value for c in row] for row in ws.iter_rows()]
        # Chaque valeur dans sa propre cellule (jamais tout dans la colonne A) et nombres restés numériques.
        self.assertEqual(lignes[0], self.payload["colonnes"])
        self.assertEqual(lignes[1], ["Data analyst", "Active", 12, 25])
        self.assertEqual(lignes[2], ["Comptable", "Clôturée", 4, 0])

    def test_plafond_de_taille(self):
        from jobs.views.recruteur_exports import EXPORT_GRAPHIQUE_MAX_COLONNES, EXPORT_GRAPHIQUE_MAX_LIGNES
        self.client.force_authenticate(user=self.rh)
        limite = {"colonnes": ["A"], "lignes": [[1]] * EXPORT_GRAPHIQUE_MAX_LIGNES}
        self.assertEqual(self.client.post(self.url, limite, format="json").status_code, status.HTTP_200_OK)
        trop_de_lignes = {"colonnes": ["A"], "lignes": [[1]] * (EXPORT_GRAPHIQUE_MAX_LIGNES + 1)}
        self.assertEqual(self.client.post(self.url, trop_de_lignes, format="json").status_code, status.HTTP_400_BAD_REQUEST)
        trop_de_colonnes = {"colonnes": ["A"] * (EXPORT_GRAPHIQUE_MAX_COLONNES + 1), "lignes": []}
        self.assertEqual(self.client.post(self.url, trop_de_colonnes, format="json").status_code, status.HTTP_400_BAD_REQUEST)

    def test_format_invalide_renvoie_400(self):
        self.client.force_authenticate(user=self.rh)
        for corps in ({"titre": "x"}, {"colonnes": [], "lignes": []}, {"colonnes": ["A"], "lignes": "pas une liste"}):
            with self.subTest(corps=corps):
                response = self.client.post(self.url, corps, format="json")
                self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
