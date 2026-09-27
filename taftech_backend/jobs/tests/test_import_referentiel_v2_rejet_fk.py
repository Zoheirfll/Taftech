# -*- coding: utf-8 -*-
"""RECETTE T17 — une ligne METIER_FORMATIONS référençant un metier_id/specialite_id
introuvable doit rejeter le lot ENTIER avant publication, avec une erreur explicite citant
le numéro de ligne — jamais un import silencieux qui ignore juste la ligne fautive."""
import os
import tempfile

import openpyxl
from django.conf import settings
from django.core.management import call_command
from django.core.management.base import CommandError
from django.test import TestCase

from jobs.models import Metier


def _classeur_metiers_minimal(metier_formation_row):
    """Classeur métiers minimal (les sheets requises avant METIER_FORMATIONS, toutes vides
    sauf en-têtes) + une seule ligne METIER_FORMATIONS fournie par l'appelant."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    # Table partagée par les deux classeurs (IMPORT étape 19), importée avant tout le reste —
    # une fixture synthétique sans cet onglet fait échouer l'import sur un KeyError avant même
    # d'atteindre METIER_FORMATIONS (le comportement réellement testé par ce fichier).
    ws = wb.create_sheet("SOURCES")
    ws.append([
        "source_id", "titre", "nature", "organisme", "fichier", "url",
        "edition_annee", "consulte_le", "sha256", "perimetre_limite",
    ])

    ws = wb.create_sheet("GROUPES")
    ws.append(["groupe_id", "type_groupe", "libelle_fr", "code_source", "parent_id"])

    ws = wb.create_sheet("METIERS")
    ws.append(["metier_id", "libelle_fr", "cle_recherche", "granularite", "groupe_principal_id",
               "code_source", "libelle_source", "selectionnable", "statut", "metier_canonique_id"])

    ws = wb.create_sheet("APPELLATIONS")
    ws.append(["appellation_id", "metier_id", "libelle", "cle_recherche", "langue", "type_alias", "ambiguite", "actif"])

    ws = wb.create_sheet("COMPETENCES")
    ws.append(["competence_id", "libelle_fr", "actif"])

    ws = wb.create_sheet("ALIAS_COMPETENCES")
    ws.append(["competence_id", "libelle"])

    ws = wb.create_sheet("PROFILS_SUGGESTIONS")
    ws.append(["profil_suggestion_id", "regle_intitule_contexte", "nb_metiers"])

    ws = wb.create_sheet("METIER_COMPETENCES")
    ws.append(["lien_id", "metier_id", "competence_id", "profil_suggestion_id", "usage", "attribution_candidat_auto"])

    ws = wb.create_sheet("METIER_FORMATIONS")
    ws.append(["lien_id", "metier_id", "specialite_id", "profil_suggestion_id", "relation", "usage_score_auto", "precision"])
    ws.append(metier_formation_row)

    return wb


class ImportReferentielV2RejetFKTest(TestCase):
    def setUp(self):
        self.chemin_diplomes = os.path.join(str(settings.BASE_DIR), "TAFTECH_Diplomes_Specialites_Algerie_v2.xlsx")
        if not os.path.exists(self.chemin_diplomes):
            self.skipTest("Classeur Diplômes/Spécialités absent de cet environnement (dev uniquement).")

    def test_metier_id_introuvable_rejette_le_lot_avec_numero_de_ligne(self):
        wb = _classeur_metiers_minimal(["LIEN-1", "METIER-INEXISTANT", "SP-INEXISTANTE", None, "", 0, ""])
        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
            wb.save(f.name)
            chemin_metiers = f.name
        try:
            with self.assertRaises(CommandError) as ctx:
                call_command("import_referentiel_v2", diplomes=self.chemin_diplomes, metiers=chemin_metiers)
            message = str(ctx.exception)
            self.assertIn("ligne 2", message)  # en-tête = ligne 1, première (et seule) donnée = ligne 2
            self.assertIn("METIER-INEXISTANT", message)
            # Le lot entier doit être rejeté — rien ne doit avoir été publié (transaction.atomic).
            self.assertFalse(Metier.objects.filter(code="METIER-INEXISTANT").exists())
        finally:
            try:
                os.unlink(chemin_metiers)
            except PermissionError:
                pass  # Windows : le handle openpyxl du fichier peut rester verrouillé un instant
