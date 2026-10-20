"""Exports Excel du recruteur : candidatures reçues avec les informations professionnelles qu'il voit déjà
sur la candidature, limitées à SON entreprise, sans donnée personnelle sensible (adresse, naissance, sexe)."""
from io import BytesIO

from django.contrib.auth import get_user_model
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook
from rest_framework.test import APITestCase

from jobs.models import Candidature, OffreEmploi, ProfilEntreprise

User = get_user_model()
ADRESSE_SECRETE = "12 rue des Lilas"

SNAPSHOT = {
    "titre_professionnel": "Développeur web", "wilaya": "31 - Oran", "commune": "Bir El Djir",
    "adresse": ADRESSE_SECRETE, "date_naissance": "1995-04-12", "sexe": "HOMME",
    "diplome_ref_libelle": "Master académique", "metier": "Études et développement informatique",
    "metier_souhaite": "Gestionnaire de Paie", "salaire_souhaite": "80000 DA", "mobilite": "NATIONALE",
    "situation_actuelle": "EN_RECHERCHE", "competences": "Python, React", "langues": "Arabe:Maternelle",
    "experiences": [{"titre_poste": "Dev"}, {"titre_poste": "Dev 2"}],
}


def lire(reponse):
    feuille = load_workbook(BytesIO(reponse.content)).active
    lignes = list(feuille.iter_rows(values_only=True))
    entetes = list(lignes[0])
    return entetes, [dict(zip(entetes, ligne)) for ligne in lignes[1:]]


class ExportsRecruteurTests(APITestCase):
    def setUp(self):
        self.rh = User.objects.create_user(username="rh", email="rh@test.dz", role="RECRUTEUR")
        self.entreprise = ProfilEntreprise.objects.create(
            user=self.rh, nom_entreprise="TestCo", est_approuvee=True, registre_commerce="RC-1")
        self.offre = OffreEmploi.objects.create(entreprise=self.entreprise, titre="Ingénieur logiciel")
        self.cand = User.objects.create_user(
            username="cand", email="cand@test.dz", role="CANDIDAT", first_name="Karim", last_name="Ali", telephone="0550000001")
        self.candidature = Candidature.objects.create(
            offre=self.offre, candidat=self.cand, statut="RETENU", score_matching=77, profil_snapshot=SNAPSHOT,
            date_entretien=timezone.now(),
        )
        Candidature.objects.create(
            offre=self.offre, est_rapide=True, nom_rapide="Benali", prenom_rapide="Sara",
            email_rapide="sara@test.dz", telephone_rapide="0660000002", statut="NOUVEAU")

        self.autre_rh = User.objects.create_user(username="rh2", email="rh2@test.dz", role="RECRUTEUR")
        self.autre_entreprise = ProfilEntreprise.objects.create(
            user=self.autre_rh, nom_entreprise="Autre", est_approuvee=True, registre_commerce="RC-2")
        self.autre_offre = OffreEmploi.objects.create(entreprise=self.autre_entreprise, titre="Secrète")
        Candidature.objects.create(offre=self.autre_offre, candidat=self.cand, statut="NOUVEAU")
        self.client.force_authenticate(user=self.rh)

    def test_export_global_colonnes_a_jour_et_valeurs(self):
        entetes, lignes = lire(self.client.get(reverse('export-candidatures-excel')))
        for colonne in ('Candidat', 'Email', 'Téléphone', 'Offre', 'Titre professionnel', 'Wilaya', 'Commune', 'Métier',
                        'Métier souhaité', 'Diplôme', 'Prétention salariale', 'Mobilité', 'Situation actuelle',
                        'Compétences', 'Langues', 'Nb expériences', 'Source', 'Candidature rapide', "Date d'entretien",
                        'Statut', 'Score matching IA (%)', 'Note globale (/20)'):
            self.assertIn(colonne, entetes)
        ligne = next(l for l in lignes if l['Email'] == 'cand@test.dz')
        self.assertEqual(ligne['Titre professionnel'], 'Développeur web')
        self.assertEqual(ligne['Wilaya'], '31 - Oran')
        self.assertEqual(ligne['Métier'], 'Études et développement informatique')
        self.assertEqual(ligne['Compétences'], 'Python, React')
        self.assertEqual(ligne['Nb expériences'], 2)
        self.assertEqual(ligne['Score matching IA (%)'], 77)
        self.assertEqual(ligne['Candidature rapide'], 'Non')

    def test_candidature_rapide_sans_profil(self):
        _, lignes = lire(self.client.get(reverse('export-candidatures-excel')))
        rapide = next(l for l in lignes if l['Email'] == 'sara@test.dz')
        self.assertEqual((rapide['Candidat'], rapide['Candidature rapide']), ('Benali Sara', 'Oui'))
        self.assertIsNone(rapide['Titre professionnel'])

    def test_uniquement_les_candidatures_de_son_entreprise(self):
        _, lignes = lire(self.client.get(reverse('export-candidatures-excel')))
        self.assertEqual({l['Offre'] for l in lignes}, {'Ingénieur logiciel'})
        self.assertEqual(len(lignes), 2)

    def test_jamais_de_donnee_personnelle_sensible(self):
        reponse = self.client.get(reverse('export-candidatures-excel'))
        entetes, _ = lire(reponse)
        for interdite in ('Adresse', 'Date de naissance', 'Sexe', 'NIN'):
            self.assertNotIn(interdite, entetes)
        cellules = [str(c.value) for ligne in load_workbook(BytesIO(reponse.content)).active.iter_rows() for c in ligne]
        for valeur in (ADRESSE_SECRETE, '1995-04-12', 'HOMME'):
            self.assertNotIn(valeur, cellules)

    def test_export_d_une_offre(self):
        entetes, lignes = lire(self.client.get(reverse('export-candidatures-offre-excel', args=[self.offre.id])))
        self.assertNotIn('Offre', entetes)  # inutile : une seule offre
        self.assertIn('Titre professionnel', entetes)
        self.assertEqual(len(lignes), 2)

    def test_l_offre_d_une_autre_entreprise_est_introuvable(self):
        reponse = self.client.get(reverse('export-candidatures-offre-excel', args=[self.autre_offre.id]))
        self.assertEqual(reponse.status_code, 404)

    def test_aucune_formule_excel_executable(self):
        Candidature.objects.filter(pk=self.candidature.pk).update(
            profil_snapshot={**SNAPSHOT, "titre_professionnel": '=HYPERLINK("http://evil.test","x")'})
        feuille = load_workbook(BytesIO(self.client.get(reverse('export-candidatures-excel')).content)).active
        self.assertTrue(all(c.data_type != 'f' for ligne in feuille.iter_rows() for c in ligne))
