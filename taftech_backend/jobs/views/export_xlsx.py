"""Génération de classeurs Excel (.xlsx) partagée par tous les exports de la plateforme.

Plus aucun export CSV : le séparateur « ; » + BOM s'ouvrait avec toutes les données dans
une seule colonne sur smartphone (correction client n°16).
"""
from decimal import Decimal

from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

XLSX_CONTENT_TYPE = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
LARGEUR_COLONNE_MAX = 50
LONGUEUR_MAX_ONGLET = 31  # limite Excel pour un nom d'onglet


def _valeur_cellule(valeur):
    """Types simples conservés (les nombres restent numériques, `Decimal` compris : un score de 77.00
    s'exporte en nombre, pas en texte) ; le reste en texte."""
    if valeur is None or isinstance(valeur, (int, float, str)):
        return valeur
    if isinstance(valeur, Decimal):
        return float(valeur)
    return str(valeur)


def neutraliser_formules(ws):
    """Sécurité : openpyxl interprète toute chaîne commençant par « = » comme une formule.

    Les textes viennent d'utilisateurs (nom, entreprise, titre d'offre, candidature rapide
    anonyme) : on force le type « texte » sur chaque chaîne pour qu'une valeur comme
    `=HYPERLINK(...)` ne s'exécute jamais à l'ouverture du fichier par un admin/recruteur.
    """
    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, str):
                cell.data_type = 's'


def _remplir_onglet(ws, titre_onglet, colonnes, lignes):
    """Remplit une feuille : en-têtes en gras, formules neutralisées, largeurs, volet figé."""
    ws.title = (titre_onglet or "Export")[:LONGUEUR_MAX_ONGLET]

    ws.append([str(c) for c in colonnes])
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(vertical="center")

    for ligne in lignes:
        if not isinstance(ligne, (list, tuple)):
            continue
        ws.append([_valeur_cellule(v) for v in ligne])
    neutraliser_formules(ws)

    for col_cells in ws.columns:
        longueur = max((len(str(c.value)) for c in col_cells if c.value is not None), default=10)
        ws.column_dimensions[col_cells[0].column_letter].width = min(longueur + 2, LARGEUR_COLONNE_MAX)
    ws.freeze_panes = "A2"


def _reponse_classeur(wb, nom_fichier):
    response = HttpResponse(content_type=XLSX_CONTENT_TYPE)
    response['Content-Disposition'] = f'attachment; filename="{nom_fichier}.xlsx"'
    wb.save(response)
    return response


def reponse_xlsx(nom_fichier, titre_onglet, colonnes, lignes):
    """Construit la réponse HTTP d'un classeur à une feuille."""
    wb = Workbook()
    _remplir_onglet(wb.active, titre_onglet, colonnes, lignes)
    return _reponse_classeur(wb, nom_fichier)


def reponse_xlsx_multi(nom_fichier, onglets):
    """Classeur à plusieurs feuilles ; `onglets` = [(titre, colonnes, lignes), ...]."""
    wb = Workbook()
    wb.remove(wb.active)
    for titre, colonnes, lignes in onglets:
        _remplir_onglet(wb.create_sheet(), titre, colonnes, lignes)
    return _reponse_classeur(wb, nom_fichier)
